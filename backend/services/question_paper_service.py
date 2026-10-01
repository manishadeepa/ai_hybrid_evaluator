"""Validate complete workbooks before publishing a test's question paper."""
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import math
import ntpath
import os
from pathlib import Path, PureWindowsPath
import re
from uuid import uuid4

from backend.repositories.question_paper_repository import QuestionPaperRepository
from backend.services.test_service import TestService
from backend.services.test_dependency_service import TestDependencyService
from backend.services.question_type_schema import validate_test_type, validate_question_type, detect_paper_type


COLUMNS = ("Question No", "Question", "Marks", "CO", "LO", "Knowledge Type", "Domain", "RBT level", "Answer Key")


class QuestionPaperDependencyError(ValueError):
    """A valid upload is blocked by existing activity, not its workbook format."""


class QuestionPaperService:
    def __init__(self, tests=None, upload_dir=None):
        self.tests = tests or TestService()
        self.repository = QuestionPaperRepository(self.tests.repository)
        self.upload_dir = Path(upload_dir or self.tests.dependencies.upload_dir)

    @staticmethod
    def _filename(name):
        if (not isinstance(name, str) or not name.strip() or name != name.strip()
                or PureWindowsPath(name).name != name or any(c in name for c in '/\\:<>"|?*')
                or any(ord(c) < 32 for c in name) or name.endswith((".", " "))
                or (ntpath.isreserved(name) if hasattr(ntpath, "isreserved") else PureWindowsPath(name).is_reserved())):
            raise ValueError("Invalid question-paper filename; use a plain filename.")
        if Path(name).suffix.lower() not in (".xlsx", ".xls"):
            raise ValueError("Unsupported question-paper type; upload an Excel .xlsx or .xls workbook.")
        return name

    def _path(self, name):
        self._filename(name)
        root = self.upload_dir.resolve()
        path = root / name
        if path.resolve().parent != root:
            raise ValueError("Question-paper path must stay within the upload directory.")
        return path

    @classmethod
    def validate_workbook(cls, data, filename, test_type=None):
        if test_type is not None:
            validate_test_type(test_type)
        cls._filename(filename)
        if not isinstance(data, bytes) or not data:
            raise ValueError("Question paper is empty or unreadable.")
        try:
            if Path(filename).suffix.lower() == ".xlsx":
                from openpyxl import load_workbook
                book = load_workbook(BytesIO(data), read_only=True, data_only=False)
                try:
                    sheet = book.worksheets[0]
                    rows = []
                    for row in sheet.iter_rows():
                        if any(c.data_type in ("f", "e") for c in row):
                            raise ValueError("Question paper contains a formula or Excel error; use literal values.")
                        rows.append([c.value for c in row])
                finally:
                    book.close()
            else:
                import pandas as pd
                frame = pd.read_excel(BytesIO(data), header=None, dtype=object)
                rows = frame.where(frame.notna(), None).values.tolist()
        except ImportError as exc:
            raise ValueError("Unable to read .xls workbook; save it as .xlsx and upload again.") from exc
        except ValueError as exc:
            if str(exc).startswith("Question paper contains"):
                raise
            raise ValueError("Unable to read question paper; upload a valid Excel workbook.") from exc
        except Exception as exc:
            raise ValueError("Unable to read question paper; upload a valid Excel workbook.") from exc
        # Spreadsheet formatting may extend beyond the final question.
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        if not rows:
            raise ValueError("Question paper contains no questions.")
        headers = rows[0]
        present = [h for h in headers if h is not None]
        if any(not isinstance(h, str) for h in present):
            raise ValueError("Question-paper column names must be text.")
        if len(present) != len(set(present)):
            raise ValueError("Question paper contains duplicate column names.")
        missing = [c for c in COLUMNS if c not in headers]
        if missing:
            raise ValueError("Question paper is missing required columns: " + ", ".join(missing))
        if len(rows) == 1:
            raise ValueError("Question paper contains no questions.")
        indices = {c: headers.index(c) for c in COLUMNS}
        questions, numbers, texts = [], set(), set()
        def scalar(value):
            return isinstance(value, (str, int, float)) and not isinstance(value, bool) and (not isinstance(value, float) or math.isfinite(value))
        for row_number, values in enumerate(rows[1:], start=2):
            row = {c: values[i] if i < len(values) else None for c, i in indices.items()}
            raw = row["Question No"]
            match = re.fullmatch(r"Q?([1-9][0-9]*)(?:\.0)?", str(raw).strip()) if scalar(raw) else None
            if not match:
                raise ValueError(f"Row {row_number} has a missing or invalid question number; use Q1 or 1.")
            number = int(match.group(1))
            label = f"Question Q{number}"
            if number in numbers:
                raise ValueError(f"Duplicate question number: Q{number}.")
            numbers.add(number)
            text = row["Question"]
            if not isinstance(text, str) or not text.strip():
                raise ValueError(f"{label} has missing question text.")
            # The unchanged AI loader indexes by stripped question text.
            if text.strip() in texts:
                raise ValueError(f"{label} duplicates another question's text.")
            texts.add(text.strip())
            try:
                if isinstance(row["Marks"], bool):
                    raise ValueError()
                marks = float(row["Marks"])
            except (TypeError, ValueError):
                raise ValueError(f"{label} has invalid marks; use a nonnegative number.") from None
            if not math.isfinite(marks) or marks < 0:
                raise ValueError(f"{label} has invalid marks; use a finite nonnegative number.")
            row["Marks"] = marks
            for field, pattern in (("CO", r"(?:CO)?[1-9][0-9]*(?:\.0)?"), ("LO", r"(?:LO)?[1-9][0-9]*(?:\.[0-9]+)*")):
                value = row[field]
                if not scalar(value) or not re.fullmatch(pattern, str(value).strip()):
                    raise ValueError(f"{label} has missing or invalid {field}.")
            # No closed metadata vocabulary exists; retain the application's labels.
            for field in ("Knowledge Type", "Domain", "RBT level"):
                if not isinstance(row[field], str) or not row[field].strip():
                    raise ValueError(f"{label} has missing or invalid {field}.")
            answer = row["Answer Key"]
            if not scalar(answer) or not str(answer).strip():
                raise ValueError(f"{label} has a missing or invalid answer key.")
            if test_type is not None:
                try:
                    row.update(validate_question_type(row, test_type))
                except ValueError as exc:
                    raise ValueError(f"{label}: {exc}") from exc
            questions.append(row)
        detected = detect_paper_type(questions)
        total = sum(q["Marks"] for q in questions)
        if not math.isfinite(total):
            raise ValueError("Question-paper total marks are not finite.")
        return {"questions": questions, "question_count": len(questions), "total_marks": total,
                "test_type": detected, "schema_version": 1}

    def import_file(self, assessment_id, test_id, source):
        source = Path(source)
        self._filename(source.name)
        try:
            data = source.read_bytes()
        except OSError as exc:
            raise ValueError("Question-paper file is missing or unreadable.") from exc
        return self.import_upload(assessment_id, test_id, source.name, data)

    def _check_dependencies(self, test, *, first_upload=False):
        # A filename-only Formative 1 response cannot identify this assessment.
        # Initial association changes no existing paper; replacement/detach and
        # delete/rename remain conservative about ambiguous historical responses.
        try:
            TestDependencyService(self.tests.repository.file_path.parent, self.upload_dir).ensure_mutable(
                {**test, "question_paper": ""}, self.tests.get_assessment(test["test_id"]),
                include_unscoped_legacy_responses=not first_upload)
        except ValueError as exc:
            detail = str(exc).removeprefix("Cannot delete or rename test because ")
            raise QuestionPaperDependencyError("Cannot change question paper: " + detail) from exc

    def import_upload(self, assessment_id, test_id, filename, data):
        with self.tests._lock():
            test = self.tests.get_test(test_id, assessment_id)
            validated = self.validate_workbook(data, filename)
            self._check_dependencies(test, first_upload=not bool(test.get("question_paper")))
            self.repository.get_by_test(test_id)  # Do not overwrite corrupt existing metadata.
            self.upload_dir.mkdir(parents=True, exist_ok=True)
            destination = self._path(filename)
            if destination.exists():
                destination = self._path(f"{Path(filename).stem}_{uuid4().hex}{Path(filename).suffix}")
            paper = {"question_paper_id": uuid4().hex, "assessment_id": assessment_id, "test_id": test_id,
                     "filename": destination.name, "original_filename": filename,
                     "uploaded_at": datetime.now(timezone.utc).isoformat(), "sha256": sha256(data).hexdigest(), **validated}
            created = False
            try:
                with destination.open("xb") as output:
                    created = True
                    output.write(data)
                    output.flush()
                    os.fsync(output.fileno())
                self.repository.save_for_test(test_id, paper)
            except (OSError, ValueError):
                if created:
                    destination.unlink(missing_ok=True)
                raise
            return paper

    def get_paper(self, assessment_id, test_id):
        with self.tests._lock():
            test = self.tests.get_test(test_id, assessment_id)
            filename = test.get("question_paper")
            if not filename:
                return None
            try:
                data = self._path(filename).read_bytes()
            except OSError as exc:
                raise ValueError("Question-paper file is missing or unreadable.") from exc
            stored = self.repository.get_by_test(test_id)
            if stored:
                if stored.get("sha256") != sha256(data).hexdigest():
                    raise ValueError("Question-paper file changed outside the application; import it again.")
                if stored.get("test_type") != test.get("test_type"):
                    raise ValueError("Question-paper type differs from the test; import the paper again.")
                return stored
            # Read legacy associations without migrating or changing them.
            return {"assessment_id": assessment_id, "test_id": test_id, "filename": filename,
                    **self.validate_workbook(data, filename, test.get("test_type"))}

    def get_questions(self, assessment_id, test_id):
        paper = self.get_paper(assessment_id, test_id)
        return paper["questions"] if paper else []

    def list_papers(self, assessment_id=None):
        return [self.get_paper(t["assessment_id"], t["test_id"]) for t in self.tests.list_tests(assessment_id)
                if t.get("question_paper")]

    def remove_paper(self, assessment_id, test_id):
        with self.tests._lock():
            test = self.tests.get_test(test_id, assessment_id)
            self._check_dependencies(test)
            self.repository.get_by_test(test_id)
            self.repository.save_for_test(test_id, None)
            # Keep uploaded files, including legacy/shared files, untouched.
