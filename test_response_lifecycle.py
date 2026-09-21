"""Response lifecycle tests use isolated JSON and the real Excel writer; no Azure."""
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from backend.repositories.candidate_repository import CandidateRepository
from backend.services.candidate_management_service import CandidateManagementService
from backend.services.candidate_assignment_service import CandidateAssignmentService
from backend.services.response_lifecycle_service import ResponseLifecycleService
from backend.services.question_paper_service import QuestionPaperService, COLUMNS
from ai_hybrid_evaluator.services.candidate_response_service import find_candidate_response_file, get_latest_candidate_response


class ResponseLifecycleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        previous = Path.cwd()
        os.chdir(self.root)
        self.addCleanup(os.chdir, previous)
        self.candidates = CandidateManagementService(CandidateRepository(self.root / 'candidates.json'))
        self.assignments = CandidateAssignmentService(self.candidates)
        self.a = self.candidates.assessments.create_assessment({'name': 'A', 'tests': ['Formative 1', 'Formative 2']})
        self.b = self.candidates.assessments.create_assessment({'name': 'B', 'tests': ['Formative 1']})
        self.aid = self.a['assessment_id']
        self.tid = self.a['test_ids']['Formative 1']
        self.key = ('C1', self.aid, self.tid)
        for cid in ('C1', 'C2'):
            self.candidates.create_candidate({'candidate_id': cid, 'name': cid + ' Name', 'email': cid + '@example.com'})
            for assessment in (self.a, self.b):
                self.assignments.assign_candidate_to_assessment(cid, assessment['assessment_id'])
                for tid in assessment['test_ids'].values():
                    self.assignments.assign_candidate_to_test(cid, tid)
        self.questions = [dict(id=i, title=f'Q{i}', text=f'Question {i}', marks=2, co='CO1', lo='LO1',
                               knowledge_type='Fact', category='Cognitive', rbt_level='Remember') for i in (1, 2)]
        self.service = self.fresh()

    def fresh(self):
        return ResponseLifecycleService(self.assignments, question_loader=lambda *_: self.questions)

    def start(self):
        return self.service.start_test(*self.key)

    def assignment_status(self):
        return self.assignments.get_test_assignment('C1', self.tid)['status']

    def test_start_assignment_and_timestamps(self):
        self.assertEqual(self.service.get_status(*self.key), 'Assigned')
        value = self.start()
        self.assertEqual(self.assignment_status(), 'In Progress')
        self.assertEqual(value['started_at'], self.assignments.get_test_assignment('C1', self.tid)['started_at'])
        self.assertEqual(value['answers'], {})
        self.assertTrue(value['created_at'])

    def test_duplicate_start_and_resume_preserve_session(self):
        old = self.start()
        self.service.save_answer(*self.key, 1, 'A')
        self.service.mark_for_review(*self.key, 1)
        self.service.record_violation(*self.key)
        value = self.service.start_test(*self.key)
        self.assertEqual(value, self.fresh().resume_test(*self.key))
        self.assertEqual(value['started_at'], old['started_at'])
        self.assertEqual(value['answers'], {'1': 'A'})
        self.assertEqual(value['marked_for_review'], [1])
        self.assertEqual(value['violation_count'], 1)
        self.assertEqual(len(self.service.repository.get_all()), 1)

    def test_save_updates_only_one_answer(self):
        self.start()
        for question, answer in ((1, 'A'), (2, 'B'), (1, 'Changed')):
            self.service.save_answer(*self.key, question, answer)
        self.assertEqual(self.fresh().get_answers(*self.key), {'1': 'Changed', '2': 'B'})

    def test_empty_answer_requires_explicit_clear(self):
        self.start()
        self.service.save_answer(*self.key, 1, 'A')
        self.service.save_answer(*self.key, 1, '  ')
        self.assertEqual(self.service.get_answers(*self.key), {'1': 'A'})
        self.service.save_answer(*self.key, 1, '', clear=True)
        self.assertEqual(self.service.get_answers(*self.key), {'1': ''})

    def test_review_is_unique_and_removes_only_requested(self):
        self.start()
        for q in (1, 1, 2): self.service.mark_for_review(*self.key, q)
        self.assertEqual(self.service.mark_for_review(*self.key, 1, False)['marked_for_review'], [2])

    def test_violation_persistence_and_monotonicity(self):
        self.start()
        self.service.record_violation(*self.key)
        self.assertEqual(self.fresh().resume_test(*self.key)['violation_count'], 1)
        for invalid in (0, -1, True, '2'):
            with self.assertRaises(ValueError): self.service.update_violation_count(*self.key, invalid)

    def test_violation_cap_disqualifies(self):
        self.start()
        value = self.service.update_violation_count(*self.key, 99)
        self.assertEqual(value['violation_count'], 3)
        self.assertEqual(value['status'], 'Disqualified')
        self.assertEqual(self.assignment_status(), 'Disqualified')
        self.assertTrue(value['terminated_at'])
        self.assertTrue(value['submission_receipt'].startswith('DQ-'))

    def test_disqualification_preserves_answers_and_does_not_write_excel(self):
        self.start()
        self.service.save_answer(*self.key, 1, 'A')
        with patch.object(self.service, 'excel_writer') as writer:
            value = self.service.disqualify(*self.key)
        writer.assert_not_called()
        self.assertEqual(value['answers'], {'1': 'A'})
        self.assertEqual(value['response_file'], '')

    def assert_terminal_guards(self):
        actions = [lambda: self.start(), lambda: self.service.resume_test(*self.key),
                   lambda: self.service.save_answer(*self.key, 1, 'bad'),
                   lambda: self.service.mark_for_review(*self.key, 1),
                   lambda: self.service.record_violation(*self.key),
                   lambda: self.service.submit(*self.key), lambda: self.service.disqualify(*self.key)]
        before = self.service.repository.file_path.read_bytes()
        for action in actions:
            with self.assertRaises(ValueError): action()
        self.assertEqual(before, self.service.repository.file_path.read_bytes())

    def test_submitted_is_terminal(self):
        self.start(); self.service.submit(*self.key)
        self.assert_terminal_guards()

    def test_disqualified_is_terminal(self):
        self.start(); self.service.disqualify(*self.key)
        self.assert_terminal_guards()

    def test_normal_submission_real_excel_and_lookup(self):
        self.start()
        value = self.service.submit(*self.key, answers={'1': 'Actual answer'})
        self.assertEqual(self.assignment_status(), 'Submitted')
        self.assertEqual(value['status'], 'Submitted')
        self.assertTrue(value['submitted_at'])
        self.assertTrue(value['submission_receipt'].startswith('SUB-'))
        lookup = dict(candidate_id='C1', assessment_name='A', test_name='Formative 1', require_assessment_scope=True)
        self.assertEqual(find_candidate_response_file(**lookup).resolve(), Path(value['response_file']))
        loaded = get_latest_candidate_response(**lookup)
        self.assertEqual([r['response'] for r in loaded['responses']], ['Actual answer', ''])
        self.assertEqual(loaded['responses'][0]['max_marks'], '2')
        book = load_workbook(value['response_file'], read_only=True)
        try:
            self.assertEqual(book.sheetnames, ['Responses', 'Submission'])
            self.assertEqual(list(book['Submission'].values)[1], ('C1', 'C1 Name', 'A', 'Formative 1'))
        finally: book.close()

    def test_timer_expiry_uses_same_artifact_and_terminal_status(self):
        self.start()
        value = self.service.handle_time_expired(*self.key, answers={'2': 'Timed answer'})
        self.assertEqual(value['submission_reason'], 'time_expired')
        self.assertTrue(Path(value['response_file']).is_file())
        self.assertEqual(value['answers'], {'2': 'Timed answer'})
        self.assert_terminal_guards()

    def test_receipts_are_persisted_and_unique(self):
        receipts = []
        for cid in ('C1', 'C2'):
            key = (cid, self.aid, self.tid)
            self.service.start_test(*key)
            value = self.service.submit(*key)
            self.assertEqual(self.fresh().get_submission_metadata(*key)['submission_receipt'], value['submission_receipt'])
            receipts.append(value['submission_receipt'])
        self.assertNotEqual(*receipts)

    def test_candidate_assessment_test_isolation(self):
        self.start(); self.service.save_answer(*self.key, 1, 'Private')
        keys = [('C2', self.aid, self.tid), ('C1', self.b['assessment_id'], self.b['test_ids']['Formative 1']),
                ('C1', self.aid, self.a['test_ids']['Formative 2'])]
        for key in keys:
            self.assertIsNone(self.service.get_response(*key))
            self.service.start_test(*key)
            self.assertEqual(self.service.get_answers(*key), {})
        with self.assertRaises(ValueError): self.service.get_response('C1', self.b['assessment_id'], self.tid)
        self.assertEqual(self.service.get_answers(*self.key), {'1': 'Private'})

    def test_unassigned_and_missing_session_rejected(self):
        with self.assertRaises(ValueError): self.service.save_answer(*self.key, 1, 'A')
        with self.assertRaises(ValueError): self.service.resume_test(*self.key)
        self.candidates.create_candidate({'candidate_id': 'C3', 'name': 'Three', 'email': 'three@example.com'})
        with self.assertRaises(ValueError): self.service.start_test('C3', self.aid, self.tid)
        self.assertFalse(self.service.repository.file_path.exists())

    def test_invalid_answer_and_question_do_not_mutate(self):
        self.start(); before = self.service.repository.file_path.read_bytes()
        for q, answer in ((99, 'A'), (True, 'A'), (1.5, 'A'), (1, None)):
            with self.assertRaises(ValueError): self.service.save_answer(*self.key, q, answer)
        with self.assertRaises(ValueError): self.service.submit(*self.key, answers={'1': 'A'}, reason='invalid')
        self.assertEqual(before, self.service.repository.file_path.read_bytes())

    def test_corrupt_json_fails_closed(self):
        self.start()
        for content in ('broken', '{}', ''):
            self.service.repository.file_path.write_text(content, encoding='utf-8')
            with self.assertRaises(ValueError): self.service.save_answer(*self.key, 1, 'A')
            self.assertEqual(self.service.repository.file_path.read_text(), content)

    def test_duplicate_or_wrong_question_stored_data_fails_closed(self):
        record = self.start()
        for rows in ([record, record], [{**record, 'answers': {'99': 'Invalid'}}], [{**record, '_pending_status': 'Submitted'}]):
            content = json.dumps(rows)
            self.service.repository.file_path.write_text(content)
            with self.assertRaises(ValueError): self.service.get_response(*self.key)
            self.assertEqual(self.service.repository.file_path.read_text(), content)

    def test_excel_failure_preserves_draft_and_assignment(self):
        self.start()
        with patch.object(self.service, 'excel_writer', side_effect=OSError('disk failure')):
            with self.assertRaisesRegex(ValueError, 'has not been submitted'):
                self.service.submit(*self.key, answers={'1': 'Final draft'})
        self.assertEqual(self.assignment_status(), 'In Progress')
        value = self.fresh().resume_test(*self.key)
        self.assertEqual(value['answers'], {'1': 'Final draft'})
        self.assertEqual(value['submission_receipt'], '')
        self.assertIsNone(value['submitted_at'])

    def test_start_persistence_failure_does_not_change_assignment(self):
        with patch.object(self.service.repository, 'save', side_effect=OSError('disk failure')):
            with self.assertRaises(OSError): self.start()
        self.assertEqual(self.assignment_status(), 'Assigned')

    def test_pending_start_recovers_after_assignment_write_failure(self):
        with patch.object(self.assignments, 'update_test_status', side_effect=OSError('disk failure')):
            with self.assertRaises(OSError): self.start()
        self.assertEqual(self.assignment_status(), 'Assigned')
        self.assertEqual(self.fresh().resume_test(*self.key)['status'], 'In Progress')
        self.assertEqual(self.assignment_status(), 'In Progress')

    def test_pending_submission_recovers_without_duplicate_excel(self):
        self.start()
        with patch.object(self.assignments, 'update_test_status', side_effect=OSError('disk failure')):
            with self.assertRaises(OSError): self.service.submit(*self.key)
        self.assertEqual(self.assignment_status(), 'In Progress')
        value = self.fresh().get_response(*self.key)
        self.assertEqual(value['status'], 'Submitted')
        self.assertEqual(self.assignment_status(), 'Submitted')
        self.assertEqual(len(list(Path('uploaded_files/candidate_responses').glob('*.xlsx'))), 1)

    def test_final_response_write_failure_is_recoverable(self):
        self.start()
        save = self.service.repository.save
        def fail_final(record):
            if record['status'] == 'Submitted' and '_pending_status' not in record:
                raise OSError('final write failure')
            return save(record)
        with patch.object(self.service.repository, 'save', side_effect=fail_final):
            with self.assertRaises(OSError): self.service.submit(*self.key)
        self.assertEqual(self.assignment_status(), 'Submitted')
        self.assertNotIn('_pending_status', self.fresh().get_response(*self.key))

    def test_failed_pending_write_removes_only_new_workbook(self):
        self.start()
        directory = Path('uploaded_files/candidate_responses'); directory.mkdir(parents=True)
        old = directory / 'existing.xlsx'; old.write_bytes(b'untouched')
        with patch.object(self.service.repository, 'save', side_effect=OSError('disk failure')):
            with self.assertRaises(OSError): self.service.submit(*self.key)
        self.assertEqual(list(directory.iterdir()), [old])
        self.assertEqual(old.read_bytes(), b'untouched')
        self.assertEqual(self.assignment_status(), 'In Progress')

    def test_missing_pending_workbook_cannot_finalize(self):
        self.start()
        with patch.object(self.assignments, 'update_test_status', side_effect=OSError('disk failure')):
            with self.assertRaises(OSError): self.service.submit(*self.key)
        record = self.service.repository.get(*self.key)
        Path(record['response_file']).unlink()
        with self.assertRaisesRegex(ValueError, 'workbook is missing'): self.fresh().get_response(*self.key)
        self.assertEqual(self.assignment_status(), 'In Progress')

    def test_concurrent_start_and_answers_do_not_lose_updates(self):
        with ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(lambda _: self.fresh().start_test(*self.key), (1, 2)))
            list(pool.map(lambda q: self.fresh().save_answer(*self.key, q, f'A{q}'), (1, 2)))
        self.assertEqual(len(self.service.repository.get_all()), 1)
        self.assertEqual(self.service.get_answers(*self.key), {'1': 'A1', '2': 'A2'})

    def test_question_snapshot_is_stable_and_excludes_answer_keys(self):
        self.questions[0]['Answer Key'] = 'secret rubric'
        self.start()
        self.questions[0]['text'] = 'Changed after start'
        value = self.fresh().get_latest_response(*self.key)
        self.assertEqual(value['questions'][0]['text'], 'Question 1')
        self.assertNotIn('Answer Key', value['questions'][0])

    def test_default_loader_uses_assessment_scoped_question_paper(self):
        book = Workbook(); sheet = book.active; sheet.append(list(COLUMNS))
        sheet.append(['Q1', 'Real question', 5, 'CO1', 'LO1', 'Fact', 'Cognitive', 'Remember', 'Answer key'])
        data = BytesIO(); book.save(data); book.close()
        assessment = self.candidates.assessments.create_assessment({'name': 'Paper test', 'tests': ['Formative 1']})
        aid, tid = assessment['assessment_id'], assessment['test_ids']['Formative 1']
        QuestionPaperService(self.candidates.assessments.tests).import_upload(aid, tid, 'paper.xlsx', data.getvalue())
        self.assignments.assign_candidate_to_assessment('C1', aid)
        self.assignments.assign_candidate_to_test('C1', tid)
        key = ('C1', aid, tid)
        service = ResponseLifecycleService(self.assignments)
        value = service.start_test(*key)
        self.assertEqual(value['questions'][0]['text'], 'Real question')
        self.assertEqual(value['questions'][0]['id'], 1)
        self.assertNotIn('Answer Key', value['questions'][0])
        self.assertTrue(Path(service.submit(*key)['response_file']).is_file())

    def test_assignment_response_mismatch_fails_closed(self):
        self.start()
        self.assignments.update_test_status('C1', self.tid, 'Submitted')
        with self.assertRaisesRegex(ValueError, 'status differ'): self.service.save_answer(*self.key, 1, 'A')

    def test_candidate_rename_is_reflected_in_submission_metadata(self):
        self.start()
        self.candidates.update_candidate('C1', {'name': 'Renamed Candidate'})
        value = self.service.submit(*self.key)
        self.assertEqual(value['candidate_name'], 'Renamed Candidate')
        self.assertTrue(Path(value['response_file']).name.startswith('Renamed_Candidate_'))


if __name__ == '__main__':
    unittest.main()
