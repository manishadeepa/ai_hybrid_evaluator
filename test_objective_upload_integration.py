"""Upload regressions: actual workbook, ambiguous legacy response, real Reflex event."""
import json
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook
from backend.repositories.candidate_repository import CandidateRepository
from backend.services.candidate_management_service import CandidateManagementService
from backend.services.candidate_assignment_service import CandidateAssignmentService
from backend.services.response_lifecycle_service import ResponseLifecycleService
from backend.services.question_paper_service import QuestionPaperService, QuestionPaperDependencyError, COLUMNS
from backend.services.question_type_schema import objective_question
from test_question_paper_management import module, rx, AdminState

ACTUAL = Path(__file__).parent/'uploaded_files'/'HoQ_Objective_Questions.xlsx'


def objective_workbook(question=None, answer='A) Option A'):
    book=Workbook();book.active.title='Objective Questions';book.active.append(COLUMNS)
    for i in range(1,15):
        text=question or f'Question stem {i}\nA) Option A\nB) Option B\nC) Option C\nD) Option D'
        if question:text=f'{i}: '+text
        book.active.append([f'Q{i}',text,2,1,1.1,'Fact','Cognitive','Remember',answer])
    stream=BytesIO();book.save(stream);book.close();return stream.getvalue()


class ObjectiveUploadIntegrationTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        temp=tempfile.TemporaryDirectory();self.addCleanup(temp.cleanup);self.root=Path(temp.name)
        self.candidates=CandidateManagementService(CandidateRepository(self.root/'candidates.json'))
        self.assessments=self.candidates.assessments;self.tests=self.assessments.tests
        self.assessment=self.assessments.save_assessment({'name':'Target assessment','tests':['Formative 1'],
            'facilitator_ids':['F001'],'assigned_candidates':[]})
        self.aid=self.assessment['assessment_id'];self.tid=self.assessment['test_ids']['Formative 1']
        self.papers=QuestionPaperService(self.tests)
        state=rx.State(_reflex_internal_init=True)
        self.fac=state.get_substate(tuple(module.FacilitatorState.get_full_name().split('.')))
        self.admin=state.get_substate(tuple(AdminState.get_full_name().split('.')))
        self.admin._apply_assessment_records(self.assessments.load_assessments())
        self.fac.selected_assessment_name=self.assessment['name'];self.fac.selected_test_name='Formative 1'
        for change in (patch.object(module,'AssessmentService',return_value=self.assessments),
                       patch.object(rx,'get_upload_dir',return_value=self.papers.upload_dir)):
            change.start();self.addCleanup(change.stop)

    def legacy(self, scoped=False, assessment=None):
        folder=self.papers.upload_dir/'candidate_responses';folder.mkdir(parents=True,exist_ok=True)
        book=Workbook();book.active.append(['Candidate Answer']);book.active.append(['Historic response'])
        if scoped:
            sheet=book.create_sheet('Submission');sheet.append(['assessment_name','test_name'])
            sheet.append([assessment or self.assessment['name'],'Formative 1'])
        path=folder/'Legacy_Formative_1_Response.xlsx';book.save(path);book.close();return path

    async def upload(self, data, filename='any-valid-name.xlsx'):
        class Upload:
            async def read(self):return data
        file=Upload();file.filename=filename
        return await self.fac.handle_upload_for_test([file])

    def assert_objective(self, count=14):
        self.assertEqual(self.fac.qp_validation_status,'success',self.fac.qp_validation_error)
        test=self.tests.get_test(self.tid,self.aid)
        self.assertEqual(test['test_type'],'objective')
        paper=self.papers.get_paper(self.aid,self.tid);self.assertEqual(paper['question_count'],count)
        self.assertEqual(set(paper['questions'][0]['options']),set('ABCD'))
        return paper

    async def test_exact_format_with_legacy_name_collision_uploads_and_snapshot_is_safe(self):
        legacy=self.legacy();before=legacy.read_bytes()
        await self.upload(objective_workbook())
        paper=self.assert_objective()
        self.assertEqual(paper['questions'][0]['correct_option'],'A')
        self.assertEqual(paper['questions'][0]['options']['A'],'Option A')
        self.assertEqual(before,legacy.read_bytes())
        self.candidates.create_candidate({'candidate_id':'C1','name':'Candidate','email':'c1@example.com'})
        assignments=CandidateAssignmentService(self.candidates)
        assignments.assign_candidate_to_assessment('C1',self.aid);assignments.assign_candidate_to_test('C1',self.tid)
        snapshot=ResponseLifecycleService(assignments).start_test('C1',self.aid,self.tid)
        self.assertEqual(snapshot['test_type'],'objective')
        for row in snapshot['questions']:
            self.assertNotIn('correct_option',row);self.assertNotIn('Answer Key',row)
            self.assertEqual(row['question_type'],'Objective');self.assertEqual(set(row['options']),set('ABCD'))

    async def test_actual_hoq_workbook_through_state_with_legacy_collision(self):
        self.legacy()
        await self.upload(ACTUAL.read_bytes(),ACTUAL.name)
        paper=self.assert_objective()
        self.assertEqual(paper['total_marks'],40)
        self.assertEqual(paper['questions'][0]['correct_option'],'A')
        self.assertEqual(paper['questions'][0]['options']['A'],'To translate customer requirements into technical requirements')

    async def test_subjective_first_upload_still_works_with_legacy_collision(self):
        self.legacy();await self.upload(objective_workbook('Explain the method','Reference answer'))
        self.assertEqual(self.fac.qp_validation_status,'success',self.fac.qp_validation_error)
        self.assertEqual(self.tests.get_test(self.tid)['test_type'],'subjective')

    async def test_replacement_and_detach_keep_legacy_protection(self):
        await self.upload(objective_workbook());self.legacy()
        before=self.tests.repository.file_path.read_bytes()
        event=await self.upload(objective_workbook())
        self.assertEqual(self.fac.qp_validation_status,'error')
        self.assertFalse(self.fac.qp_validation_popup_open)
        self.assertIn('legacy response',self.fac.qp_validation_error)
        self.assertIn('legacy response',str(event))
        self.assertNotIn('delete or rename',self.fac.qp_validation_error)
        self.assertEqual(before,self.tests.repository.file_path.read_bytes())
        with self.assertRaises(QuestionPaperDependencyError):self.papers.remove_paper(self.aid,self.tid)

    async def test_known_submission_blocks_first_upload(self):
        self.legacy(scoped=True)
        await self.upload(objective_workbook())
        self.assertEqual(self.fac.qp_validation_status,'error')
        self.assertIn('candidate response',self.fac.qp_validation_error)
        self.assertFalse(self.fac.qp_validation_popup_open)
        self.assertNotIn('test_type',self.tests.get_test(self.tid))

    async def test_foreign_scoped_response_does_not_block(self):
        self.legacy(scoped=True,assessment='Other assessment')
        await self.upload(objective_workbook());self.assert_objective()

    async def test_known_activity_blocks_first_upload(self):
        for filename,record in (
            ('ai_evaluation_runs.json',{'test_id':self.tid}),
            ('manual_evaluations.json',{'assessment_name':self.assessment['name'],'test_name':'Formative 1'}),
            ('test_candidates.json',{'test_id':self.tid,'status':'In Progress'})):
            with self.subTest(filename=filename):
                path=self.root/filename;path.write_text(json.dumps([record]),encoding='utf-8')
                try:
                    await self.upload(objective_workbook())
                    self.assertEqual(self.fac.qp_validation_status,'error')
                    self.assertFalse(self.fac.qp_validation_popup_open)
                finally:path.unlink()

    async def test_delete_and_rename_still_protected_for_legacy_collision(self):
        self.legacy()
        with self.assertRaisesRegex(ValueError,'legacy response'):self.tests.delete_test(self.aid,self.tid)
        with self.assertRaisesRegex(ValueError,'legacy response'):self.tests.update_test(self.tid,{'test_name':'Renamed'})

    async def test_invalid_answer_and_malformed_paper_remain_format_errors(self):
        for data in (objective_workbook(answer='A) Wrong'),objective_workbook('Stem\nA) A\nB) B','A')):
            with self.subTest():
                await self.upload(data)
                self.assertEqual(self.fac.qp_validation_status,'error');self.assertTrue(self.fac.qp_validation_popup_open)
                self.assertTrue(self.fac.qp_validation_error)
                self.assertNotIn('test_type',self.tests.get_test(self.tid))
        await self.upload(objective_workbook());self.assert_objective();self.assertEqual(self.fac.qp_validation_error,'')

    async def test_key_forms_are_explicit_not_silently_loosened(self):
        question='Stem\nA) Option A\nB) Option B\nC) Option C\nD) Option D'
        for key in ('A','A) Option A'):
            self.assertEqual(objective_question(question,key)['correct_option'],'A')
        for key in ('A)','A. Option A','A) Wrong'):
            with self.assertRaises(ValueError):objective_question(question,key)
        with self.assertRaises(ValueError):objective_question(question.replace(')', '.'),'A')

if __name__=='__main__':unittest.main()
